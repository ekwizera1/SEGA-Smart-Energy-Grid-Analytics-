import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

MEDIA_DIR = "extracted_media/word/media"

def style_table(table, col_widths, hdr_bg="1B365D", alt_bg="F8FAFC", border_color="CBD5E0"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>\n'
        f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="{border_color}"/>\n'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

    for i, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if i == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for j, cell in enumerate(row.cells):
            cell.width = Inches(col_widths[j])
            tcPr = cell._tc.get_or_add_tcPr()
            tcMar = parse_xml(
                f'<w:tcMar {nsdecls("w")}>\n'
                f'  <w:top w:w="120" w:type="dxa"/>\n'
                f'  <w:bottom w:w="120" w:type="dxa"/>\n'
                f'  <w:left w:w="150" w:type="dxa"/>\n'
                f'  <w:right w:w="150" w:type="dxa"/>\n'
                f'</w:tcMar>'
            )
            tcPr.append(tcMar)

            if i == 0:
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hdr_bg}"/>')
                tcPr.append(shd)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    for r in p.runs:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        r.font.size = Pt(8.5)
            else:
                if i % 2 == 0:
                    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{alt_bg}"/>')
                    tcPr.append(shd)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(1.5)
                    p.paragraph_format.space_after = Pt(1.5)
                    for r in p.runs:
                        r.font.size = Pt(8.5)
                        r.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

def add_callout(doc, title, text, bg_color="EDF2F7", border_color="1B365D"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    tcPr = cell._tc.get_or_add_tcPr()
    
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="none"/>\n'
        f'  <w:bottom w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
    tcPr.append(shd)
    
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>\n'
        f'  <w:top w:w="160" w:type="dxa"/>\n'
        f'  <w:bottom w:w="160" w:type="dxa"/>\n'
        f'  <w:left w:w="220" w:type="dxa"/>\n'
        f'  <w:right w:w="220" w:type="dxa"/>\n'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(title + "\n")
    r_title.bold = True
    r_title.font.size = Pt(9.5)
    r_title.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    
    r_text = p.add_run(text)
    r_text.font.size = Pt(9.0)
    r_text.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(4)

def add_figure(doc, img_name, caption, width_inch=5.8):
    img_path = os.path.join(MEDIA_DIR, img_name)
    if not os.path.exists(img_path):
        print(f"Warning: Image {img_path} not found")
        return
    
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(6)
    p_img.paragraph_format.space_after = Pt(2)
    r = p_img.add_run()
    r.add_picture(img_path, width=Inches(width_inch))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(1)
    p_cap.paragraph_format.space_after = Pt(8)
    r_cap = p_cap.add_run(caption)
    r_cap.font.italic = True
    r_cap.font.size = Pt(8.5)
    r_cap.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

def set_footer_pagenum(doc):
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(0)
        
        r1 = p.add_run("SEGA Technical Report | Adventist University of Central Africa (AUCA)  —  Page ")
        r1.font.size = Pt(8.0)
        r1.font.color.rgb = RGBColor(0x71, 0x80, 0x96)
        
        fld1 = OxmlElement('w:fldChar')
        fld1.set(qn('w:fldCharType'), 'begin')
        t1 = OxmlElement('w:instrText')
        t1.set(qn('xml:space'), 'preserve')
        t1.text = "PAGE"
        fld2 = OxmlElement('w:fldChar')
        fld2.set(qn('w:fldCharType'), 'separate')
        fld3 = OxmlElement('w:fldChar')
        fld3.set(qn('w:fldCharType'), 'end')
        
        p._p.append(fld1)
        p._p.append(t1)
        p._p.append(fld2)
        p._p.append(fld3)

print("Helper definitions loaded successfully.")
